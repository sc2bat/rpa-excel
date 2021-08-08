from openpyxl import load_workbook
wb = load_workbook("sample5_2.xlsx")
ws = wb.active

# 차트 만들기

from openpyxl.chart import BarChart, Reference, LineChart
# value 설정
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart() # 차트종류
# bar_chart.add_data(bar_value) # 차트 데이터 추가

# ws.add_chart(bar_chart, "E1") # 차트 넣는 위치 정의

# wb.save("sample10.xlsx")

# B1:C11 데이터 사용
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart() # 차트종류
line_chart.add_data(line_value, titles_from_data=True) # 차트 데이터 추가
line_chart.title = "coco" # 제목
line_chart.style = 20 # 정의된 스타일 이용
line_chart.y_axis.title = "score" # y축
line_chart.x_axis.title = "No" # x 축

ws.add_chart(line_chart, "E1") # 차트 넣는 위치 정의

wb.save("sample10_2.xlsx")